"""Step 1 of the accuracy check: make a blind labelling sheet.

  python -m validation.make_sample --news "C:\\path\\to\\archive.zip"
  python -m validation.make_sample --news archive.zip --extra my_real_posts.jsonl   # optional social posts

Creates  output/label_me.xlsx           <- you fill this in (the engine's answers are NOT shown)
         data/validation_engine.csv     <- the engine's answers, kept apart so you are not influenced

Then label the sheet in Excel (about 20 minutes) and run:  python -m validation.score
"""
import argparse
import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from backtest.data import load_news
from engine.config import DATA_DIR, EVENT_LABELS, OTHER_LABEL, ROOT, UNIVERSE
from engine.events import classify
from engine.impact import impact_score
from engine.sentiment import score_batch

OUTPUT_DIR = ROOT / "output"
INSTRUCTIONS = [
    "How to label (about 20 minutes for 100 rows)",
    "",
    "Column D, your_sentiment: would a reader of this text expect it to be GOOD or BAD for the company's stock?",
    "  positive = good news / upbeat (earnings beat, upgrade, price target raised, new contract)",
    "  negative = bad news / downbeat (downgrade, lawsuit, miss, recall, investigation)",
    "  neutral  = no clear effect, or just a routine notice",
    "Judge the meaning, not the individual words. Do not guess what the model will say.",
    "",
    "Column E, your_event: the MAIN topic. Pick from the dropdown; use Other when none fits.",
    "  Geopolitical, Macroeconomic, Credit Event, Merger/Acquisition, Earnings,",
    "  Product Launch, Regulatory, Leadership Change, Other",
    "",
    "Leave a cell blank if you honestly cannot tell. Save the file, then run:  python -m validation.score",
]


def build_sheet(rows: list[dict], path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Label these"
    ws.append(["id", "ticker", "text", "your_sentiment", "your_event"])
    for r in rows:
        ws.append([r["id"], r["ticker"], r["text"], None, None])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 9
    ws.column_dimensions["C"].width = 100
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 22
    for row in ws.iter_rows(min_row=2, max_col=3):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    sent = DataValidation(type="list", formula1='"positive,negative,neutral"', allow_blank=True)
    event = DataValidation(type="list", formula1='"' + ",".join(EVENT_LABELS + [OTHER_LABEL]) + '"', allow_blank=True)
    ws.add_data_validation(sent)
    ws.add_data_validation(event)
    sent.add(f"D2:D{len(rows) + 1}")
    event.add(f"E2:E{len(rows) + 1}")
    guide = wb.create_sheet("Instructions")
    for line in INSTRUCTIONS:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 110
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--news", required=True, help="Kaggle zip, folder or CSV (same one used for the backtest)")
    ap.add_argument("--n", type=int, default=100, help="how many items to label")
    ap.add_argument("--start", default="2019-01-01")
    ap.add_argument("--end", default="2019-12-31")
    ap.add_argument("--seed", type=int, default=11)
    ap.add_argument("--extra", help="optional JSONL of real posts: one {\"source\": \"social\", \"text\": \"...\"} per line")
    args = ap.parse_args()

    news = load_news(args.news, list(UNIVERSE), args.start, args.end, max_per_ticker=300)
    if news.empty:
        raise SystemExit("No headlines found. Check --news / --start / --end.")
    pool = [{"source": "news", "ticker": r.ticker, "text": r.text} for r in news.itertuples()]
    extra = []
    if args.extra:
        with open(args.extra, encoding="utf-8") as f:
            extra = [{"source": d.get("source", "social"), "ticker": d.get("ticker", ""), "text": d["text"]}
                     for d in map(json.loads, filter(str.strip, f))]
    n_extra = min(len(extra), args.n // 3)
    pd_pool = pd.DataFrame(pool).sample(n=min(len(pool), args.n - n_extra), random_state=args.seed)
    items = pd_pool.to_dict("records") + pd.DataFrame(extra).sample(n=n_extra, random_state=args.seed).to_dict("records") \
        if n_extra else pd_pool.to_dict("records")
    items = pd.DataFrame(items).sample(frac=1, random_state=args.seed).to_dict("records")
    for i, it in enumerate(items, 1):
        it["id"] = i

    print(f"Running the engine on {len(items)} items (hidden from the sheet)...")
    out = []
    for source in ("news", "social"):
        group = [it for it in items if it["source"] == source]
        if not group:
            continue
        for it, (sent, conf) in zip(group, score_batch([g["text"] for g in group], source)):
            event, event_conf, method = classify(it["text"])
            out.append({"id": it["id"], "source": source, "ticker": it["ticker"], "text": it["text"],
                        "engine_sentiment": sent, "engine_confidence": conf, "engine_event": event,
                        "engine_event_method": method, "engine_impact": impact_score(event, sent, conf, source)})
    DATA_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)
    pd.DataFrame(out).sort_values("id").to_csv(DATA_DIR / "validation_engine.csv", index=False)
    build_sheet(sorted(items, key=lambda r: r["id"]), OUTPUT_DIR / "label_me.xlsx")
    print(f"\nLabel this file:  {OUTPUT_DIR / 'label_me.xlsx'}   (open in Excel, fill columns D and E, save)")
    print("Then run:         python -m validation.score")


if __name__ == "__main__":
    main()
