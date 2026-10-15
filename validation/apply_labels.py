"""Fill the blank cells of output/label_me.xlsx from validation/draft_labels.csv (matched by id).

  python -m validation.apply_labels

Only empty cells are filled, so anything you typed yourself is kept.
Open the sheet afterwards and CHANGE every label you disagree with.
Close Excel before running this.
"""
import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT / "output" / "label_me.xlsx"
DRAFT = Path(__file__).with_name("draft_labels.csv")


def main():
    draft = {int(r["id"]): r for r in csv.DictReader(open(DRAFT, encoding="utf-8"))}
    wb = load_workbook(SHEET)
    ws = wb["Label these"]
    filled = 0
    for row in range(2, ws.max_row + 1):
        rid = ws.cell(row, 1).value
        if rid in draft:
            for col, key in ((4, "your_sentiment"), (5, "your_event")):
                if ws.cell(row, col).value in (None, ""):
                    ws.cell(row, col).value = draft[rid][key]
                    filled += 1
    wb.save(SHEET)
    print(f"Filled {filled} cells. Now open {SHEET} and fix anything you disagree with, then save.")


if __name__ == "__main__":
    main()
